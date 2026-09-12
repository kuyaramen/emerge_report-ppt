import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ==========================================
# HELPER FUNCTIONS
# ==========================================

def apply_theme(slide):
    """Sets a dark navy background."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(10, 25, 47)  # Dark navy

def format_title(shape, text, size=40):
    """Formats the title text with white color and modern font."""
    text_frame = shape.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = text
    p.font.name = 'Arial'
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
def format_body(shape, content):
    """Formats the body content with bullet points."""
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    
    for i, line in enumerate(content):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
            
        p.text = line.replace("* ", "")
        p.font.name = 'Arial'
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(230, 240, 255)
        
        if line.startswith("*") or line.startswith("-") or (line and line[0].isdigit()):
            p.level = 0
            
def add_notes(slide, notes_text):
    """Adds speaker notes to a slide."""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

# ==========================================
# PRESENTATION SETUP
# ==========================================

prs = Presentation()

# Layouts
TITLE_SLIDE_LAYOUT = prs.slide_layouts[0]
TITLE_AND_CONTENT_LAYOUT = prs.slide_layouts[1]
TWO_CONTENT_LAYOUT = prs.slide_layouts[3]

# ==========================================
# SLIDE CREATION
# ==========================================

# SLIDE 1: Title Slide
slide1 = prs.slides.add_slide(TITLE_SLIDE_LAYOUT)
apply_theme(slide1)
format_title(slide1.shapes.title, "SERVERLESS COMPUTING", size=50)
format_body(slide1.placeholders[1], [
    "Understanding Cloud-Based Application Execution\n",
    "Course: EMERGING TECHNOLOGY",
    "Instructor: Engr. Catherine M. Verallo",
    "Date: September 12, 2026\n",
    "Presented by:",
    "ABAO JOHN LOUIE B.",
    "AURE LORNA JAY",
    "KURT CHUA",
    "NIKKOS CHARLES"
])
add_notes(slide1, "Good morning, everyone. Today, we will be discussing Serverless Computing—a modern approach to how applications are built and run in the cloud. My name is [Your Name], and together with my group mates, we will explain what serverless is, how it works, and why it is so important in modern technology.")

# SLIDE 2: Learning Objectives
slide2 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide2)
format_title(slide2.shapes.title, "What We Will Learn")
format_body(slide2.placeholders[1], [
    "1. Define serverless computing.",
    "2. Explain how serverless works.",
    "3. Identify FaaS and BaaS.",
    "4. Discuss advantages and disadvantages.",
    "5. Compare cloud computing and serverless computing.",
    "6. Apply serverless concepts to a real website."
])
add_notes(slide2, "Before we dive in, let's look at our roadmap for today. We'll start by defining what cloud and serverless computing are. Then, we will look at how it works behind the scenes, the two main types, and the pros and cons. Finally, we'll apply what we've learned to a real-world example.")

# SLIDE 3: What is cloud computing?
slide3 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide3)
format_title(slide3.shapes.title, "What Is Cloud Computing?")
format_body(slide3.placeholders[1], [
    "Cloud computing is the delivery of computing services over the internet.\n",
    "Examples of services:",
    "* Servers",
    "* Storage",
    "* Databases",
    "* Networking",
    "* Software\n",
    "Instead of buying and maintaining your own computer server, you use computing resources provided by a cloud company like AWS, Azure, or Google Cloud."
])
add_notes(slide3, "To understand serverless, we first need to understand cloud computing. Put simply, cloud computing is just delivering computing services over the internet. Instead of buying a physical server and putting it in your room, you rent one from companies like Amazon, Microsoft, or Google.\n\nQuiz Question: What is cloud computing?\nAnswer: Cloud computing is the delivery of computing services over the internet.")

# SLIDE 4: What is serverless computing?
slide4 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide4)
format_title(slide4.shapes.title, "What Is Serverless Computing?")
format_body(slide4.placeholders[1], [
    "Serverless computing is a cloud computing model where developers run code without directly managing the underlying servers.\n",
    "* Servers still exist!",
    "* The cloud provider manages the infrastructure.",
    "* Developers focus entirely on writing and deploying code.",
    "* Functions run only when needed.\n",
    "IMPORTANT: Serverless does NOT mean there are no servers."
])
add_notes(slide4, "Now, what is serverless computing? It is a specific type of cloud computing where developers run code without worrying about managing the servers. An important thing to remember: Serverless does NOT mean there are no servers. Servers definitely exist in massive data centers. It just means you don't have to deal with them.\n\nQuiz Question: Why is it called serverless if servers still exist?\nAnswer: Because the cloud provider manages the underlying servers, so developers do not have to manage them directly.")

# SLIDE 5: How serverless computing works
slide5 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide5)
format_title(slide5.shapes.title, "How Serverless Computing Works")
format_body(slide5.placeholders[1], [
    "1. A user sends a request.",
    "2. The cloud receives the request.",
    "3. A serverless function runs.",
    "4. The function processes the request.",
    "5. The result is returned.",
    "6. The function execution ends.\n",
    "Workflow:",
    "User Request → Cloud Platform → Serverless Function → Process Data → Return Response"
])
add_notes(slide5, "Unlike traditional servers that are always on and waiting, serverless functions only run when something happens. A user sends a request, the cloud receives it, the code spins up, processes the data, returns the result, and then the function shuts down entirely.\n\nQuiz Question: What triggers a serverless function to run?\nAnswer: A request or event can trigger a serverless function.")

# SLIDE 6: Real-life example
slide6 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide6)
format_title(slide6.shapes.title, "Example: Get Latest News")
format_body(slide6.placeholders[1], [
    "Scenario: A user clicks 'Get Latest News'.\n",
    "1. User clicks the button.",
    "2. The request goes to the cloud.",
    "3. The serverless function runs.",
    "4. The function retrieves news from the database.",
    "5. The website displays the news.\n",
    "Flow: Website → Serverless Function → Database → Website"
])
add_notes(slide6, "Let’s look at a simple real-life example. Imagine you are on a news website and click 'Get Latest News'. That click sends a request to the cloud. A serverless function wakes up, grabs the latest news from the database, sends it back to your screen, and goes back to sleep.\n\nQuiz Question: What does the serverless function do in the news website example?\nAnswer: It processes the request and retrieves the news from the database.")

# SLIDE 7: Two main types
slide7 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide7)
format_title(slide7.shapes.title, "Two Main Types of Serverless Computing")
format_body(slide7.placeholders[1], [
    "1. Function as a Service (FaaS)",
    "Run your own code when needed.\n",
    "2. Backend as a Service (BaaS)",
    "Use ready-made backend services provided by a cloud platform."
])
add_notes(slide7, "There are two main flavors of serverless computing you need to know: FaaS and BaaS. FaaS is for when you want to run your own custom code. BaaS is when you want to use ready-made backend features like databases without building them yourself.\n\nQuiz Question: What are the two main types of serverless computing discussed in this presentation?\nAnswer: Function as a Service (FaaS) and Backend as a Service (BaaS).")

# SLIDE 8: FaaS
slide8 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide8)
format_title(slide8.shapes.title, "Function as a Service (FaaS)")
format_body(slide8.placeholders[1], [
    "FaaS allows developers to run small pieces of code in response to events or requests.\n",
    "* You write a function.",
    "* The cloud runs it when needed.",
    "* You do not manage the underlying server.",
    "* The function performs a specific task.\n",
    "Example: AWS Lambda\n",
    "function calculateTotal(price, quantity) {",
    "    return price * quantity;",
    "}"
])
add_notes(slide8, "Let's dive into FaaS. Function as a Service lets you run small pieces of code. You just write the logic, upload it, and the cloud handles the rest. In a FaaS model, code only runs when needed, and you only pay for the exact time it takes to run.\n\nQuiz Question: What does FaaS allow developers to run?\nAnswer: It allows developers to run their own functions or pieces of code when needed.")

# SLIDE 9: BaaS
slide9 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide9)
format_title(slide9.shapes.title, "Backend as a Service (BaaS)")
format_body(slide9.placeholders[1], [
    "BaaS provides ready-made backend services so developers do not need to build everything from scratch.\n",
    "Services may include:",
    "* Databases",
    "* User authentication",
    "* File storage",
    "* APIs",
    "* Notifications\n",
    "Example: Supabase"
])
add_notes(slide9, "Next is BaaS, or Backend as a Service. Building a login system or a secure database from scratch is hard. BaaS providers give you these features ready to use out of the box. You just connect your app to them.\n\nQuiz Question: What is the main purpose of BaaS?\nAnswer: To provide ready-made backend services that developers can use in their applications.")

# SLIDE 10: FaaS vs BaaS
slide10 = prs.slides.add_slide(TWO_CONTENT_LAYOUT)
apply_theme(slide10)
format_title(slide10.shapes.title, "FaaS vs BaaS")
format_body(slide10.placeholders[1], [
    "Function as a Service (FaaS):",
    "* Runs your code.",
    "* You write functions.",
    "* Focuses on code execution.",
    "* Example: AWS Lambda.\n",
    "Memory Tip: 'Run my code.'"
])
format_body(slide10.placeholders[2], [
    "Backend as a Service (BaaS):",
    "* Provides backend services.",
    "* You use ready-made features.",
    "* Focuses on backend functionality.",
    "* Example: Supabase.\n",
    "Memory Tip: 'Give me backend services.'"
])
add_notes(slide10, "To summarize the difference: FaaS is about running your specific code logic. BaaS is about giving you pre-built tools like databases. Just remember: FaaS equals 'Run my code', and BaaS equals 'Give me backend services'.\n\nQuiz Question: Which type of serverless computing focuses on running your own code?\nAnswer: FaaS — Function as a Service.")

# SLIDE 11: Advantages
slide11 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide11)
format_title(slide11.shapes.title, "Advantages of Serverless Computing")
format_body(slide11.placeholders[1], [
    "1. No direct server management (Focus on coding).",
    "2. Automatic scaling (Platform runs more instances when traffic increases).",
    "3. Usage-based pricing (Only pay when code runs).",
    "4. Faster development.",
    "5. Focus on application code."
])
add_notes(slide11, "Why do companies love serverless? First, no server maintenance. Second, it scales automatically. If your app suddenly goes viral, the platform instantly spins up more function instances to handle the traffic. Third, it's cost-effective; if nobody uses your app, your functions don't run, and you pay zero.\n\nQuiz Question: Which advantage allows serverless applications to handle changing numbers of users?\nAnswer: Automatic scaling.")

# SLIDE 12: Disadvantages
slide12 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide12)
format_title(slide12.shapes.title, "Disadvantages of Serverless Computing")
format_body(slide12.placeholders[1], [
    "1. Cold starts (A delay when a function environment needs to prepare).",
    "2. Execution time limits.",
    "3. Internet and cloud service dependency.",
    "4. Vendor lock-in.",
    "5. Costs can increase with high usage."
])
add_notes(slide12, "But it’s not perfect. A major issue is 'Cold Starts'. Waking up a sleeping function takes an extra second or two, causing a slight delay for the user. Also, you have execution limits—functions are meant for quick tasks, not hour-long processes.\n\nQuiz Question: What is a cold start?\nAnswer: A cold start is the extra startup delay that may occur when a serverless function environment needs to be prepared.")

# SLIDE 13: Traditional Server vs Serverless
slide13 = prs.slides.add_slide(TWO_CONTENT_LAYOUT)
apply_theme(slide13)
format_title(slide13.shapes.title, "Traditional Server vs Serverless")
format_body(slide13.placeholders[1], [
    "Traditional Server:",
    "* You manage the server.",
    "* Server often runs continuously.",
    "* You manage resources and configuration.",
    "* Scaling may require manual setup."
])
format_body(slide13.placeholders[2], [
    "Serverless:",
    "* Provider manages the infrastructure.",
    "* Functions run when needed.",
    "* Scaling is often automatic.",
    "* Developers focus on code."
])
add_notes(slide13, "Comparing the two directly: A traditional server is like owning a car. You have to maintain it and pay for it even when you aren't driving. Serverless is like taking a taxi. You only pay for the exact trip, and the company handles the maintenance.\n\nQuiz Question: What is the main difference between a traditional server and serverless computing?\nAnswer: In traditional server computing, the developer manages the server. In serverless computing, the cloud provider manages the underlying infrastructure.")

# SLIDE 14: Cloud vs Serverless
slide14 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide14)
format_title(slide14.shapes.title, "Cloud Computing vs Serverless Computing")
format_body(slide14.placeholders[1], [
    "Cloud computing is the broader category.",
    "Serverless computing is one model within cloud computing.\n",
    "* Cloud computing: Delivery of computing services over the internet.",
    "* Serverless computing: A cloud model where the provider manages the underlying servers for code execution and backend services."
])
add_notes(slide14, "A common point of confusion is thinking cloud and serverless are two completely different things. They aren't. Cloud computing is the massive overarching category. Serverless computing is just one specific way to use the cloud.\n\nQuiz Question: Is serverless computing a type of cloud computing?\nAnswer: Yes. Serverless computing is a cloud computing model.")

# SLIDE 15: Examples
slide15 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide15)
format_title(slide15.shapes.title, "Examples of Serverless Platforms")
format_body(slide15.placeholders[1], [
    "1. AWS Lambda",
    "Runs code in response to events (FaaS).\n",
    "2. Vercel",
    "Can host web applications and run serverless functions.\n",
    "3. Google Cloud Functions",
    "Runs backend code without requiring direct server management.\n",
    "4. Supabase",
    "Provides backend services such as databases and authentication (BaaS)."
])
add_notes(slide15, "Here are some of the biggest names in the industry. AWS Lambda and Google Cloud Functions are your classic FaaS providers. Vercel is highly popular for hosting modern websites that use serverless APIs. And Supabase is a fantastic example of a BaaS providing database services.\n\nQuiz Question: Which service is an example of a platform that runs serverless functions?\nAnswer: AWS Lambda is one example.")

# SLIDE 16: SDN PRIME
slide16 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide16)
format_title(slide16.shapes.title, "Serverless Computing in SDN PRIME")
format_body(slide16.placeholders[1], [
    "Scenario: An admin adds a new news article.\n",
    "Process:",
    "1. Admin submits a news article.",
    "2. The request goes to the backend.",
    "3. A serverless function validates and saves the data.",
    "4. The database stores the article.",
    "5. The website displays the updated news.\n",
    "Architecture:",
    "Admin → Serverless Backend Function → Supabase Database → Public Website"
])
add_notes(slide16, "Let's apply this to a practical project like SDN PRIME. If an admin wants to post news, they fill out a form. Instead of a traditional server, a Vercel serverless function catches that data, makes sure it's valid, and saves it into a Supabase database.\n\nQuiz Question: How can serverless computing help an SDN PRIME website?\nAnswer: It can run backend functions that process requests, save content, and retrieve information without requiring the project team to manage the underlying servers.")

# SLIDE 17: When to Use
slide17 = prs.slides.add_slide(TWO_CONTENT_LAYOUT)
apply_theme(slide17)
format_title(slide17.shapes.title, "When Is Serverless a Good Choice?")
format_body(slide17.placeholders[1], [
    "Serverless is useful for:",
    "* APIs.",
    "* Website backend functions.",
    "* Form submissions.",
    "* Database-triggered tasks.",
    "* Event-driven applications.",
    "* Small or changing workloads."
])
format_body(slide17.placeholders[2], [
    "Consider alternatives for:",
    "* Long-running workloads.",
    "* Highly specialized tasks.",
    "* Constant 24/7 heavy workloads (Traditional servers may be better)."
])
add_notes(slide17, "Serverless is amazing, but it isn't for everything. It is perfect for APIs, form submissions, and websites where traffic goes up and down. But if you have long-running or specialized workloads, a traditional dedicated server might be a better choice.\n\nQuiz Question: Give one example of an application task that can use serverless computing.\nAnswer: An API request, form submission, or website backend function.")

# SLIDE 18: Key Takeaways
slide18 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide18)
format_title(slide18.shapes.title, "What We Learned")
format_body(slide18.placeholders[1], [
    "1. Serverless is a cloud computing model.",
    "2. Servers still exist.",
    "3. The cloud provider manages the infrastructure.",
    "4. FaaS runs code.",
    "5. BaaS provides backend services.",
    "6. Serverless can scale automatically.",
    "7. Serverless has limitations such as cold starts and execution limits."
])
add_notes(slide18, "To wrap up our main concepts: Serverless means the cloud provider manages the servers. FaaS is for running code, BaaS is for backend services. It scales easily and saves money, but you do have to watch out for cold starts and execution limits.")

# SLIDE 19: Final Knowledge Check
slide19 = prs.slides.add_slide(TITLE_AND_CONTENT_LAYOUT)
apply_theme(slide19)
format_title(slide19.shapes.title, "Final Knowledge Check")
format_body(slide19.placeholders[1], [
    "1. What is serverless computing?",
    "2. What does FaaS stand for?",
    "3. What does BaaS stand for?",
    "4. Why is automatic scaling an advantage?",
    "5. Is serverless computing the same thing as cloud computing?"
])
add_notes(slide19, "Let's see what everyone remembers!\n\n1. What is serverless? (Answer: A cloud model where providers manage the servers.)\n2. What does FaaS stand for? (Answer: Function as a Service.)\n3. What does BaaS stand for? (Answer: Backend as a Service.)\n4. Why is automatic scaling an advantage? (Answer: It allows apps to handle changing numbers of users automatically.)\n5. Is serverless the same as cloud computing? (Answer: No, serverless is a specific model within the broader category of cloud computing.)")

# SLIDE 20: Thank You
slide20 = prs.slides.add_slide(TITLE_SLIDE_LAYOUT)
apply_theme(slide20)
format_title(slide20.shapes.title, "Thank You!", size=60)
format_body(slide20.placeholders[1], [
    "Any Questions?\n\n",
    "\"Serverless computing helps developers focus on building applications while cloud providers manage the underlying infrastructure.\""
])
add_notes(slide20, "Thank you all for listening. Serverless computing is truly the future of web development, allowing us to focus on writing great apps rather than managing hardware. Are there any questions?")

# ==========================================
# SAVE PRESENTATION
# ==========================================
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'SERVERLESS_COMPUTING_REPORT.pptx')
prs.save(out_path)
print(f"Presentation generated and saved to {out_path}")
